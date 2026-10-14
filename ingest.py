"""Document extraction, chunking, embedding, and FAISS index creation."""

from __future__ import annotations

import json
import os
import tempfile
from functools import lru_cache
from pathlib import Path
from typing import Any

import faiss
import fitz
from docx import Document
from sentence_transformers import SentenceTransformer

from utils import clean_text, unique_paths, validate_extension

CHUNK_SIZE = 3500
CHUNK_OVERLAP = 500


@lru_cache(maxsize=2)
def get_embedding_model(model_name: str) -> SentenceTransformer:
    """Load and reuse an embedding model for the current process."""
    return SentenceTransformer(model_name)


def extract_pdf(path: Path) -> list[dict[str, Any]]:
    """Extract one metadata-bearing record per PDF page."""
    records: list[dict[str, Any]] = []
    with fitz.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = clean_text(page.get_text())
            if text:
                records.append({"source": path.name, "page": page_number, "text": text})
    return records


def extract_docx(path: Path) -> list[dict[str, Any]]:
    """Extract DOCX paragraphs and tables; page numbers are unavailable."""
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    text = clean_text(text)
    return [{"source": path.name, "page": None, "text": text}] if text else []


def extract_txt(path: Path) -> list[dict[str, Any]]:
    """Extract UTF-8 text, with a small fallback for common Windows encodings."""
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="cp1252")
    text = clean_text(text)
    return [{"source": path.name, "page": None, "text": text}] if text else []


def extract_document(path: Path) -> list[dict[str, Any]]:
    """Dispatch extraction by extension and preserve source metadata."""
    validate_extension(path.name)
    extension = path.suffix.lower()
    if extension == ".pdf":
        records = extract_pdf(path)
    elif extension == ".docx":
        records = extract_docx(path)
    else:
        records = extract_txt(path)
    if not records:
        raise ValueError(f"No text could be extracted from '{path.name}'.")
    return records


def _split_text(text: str) -> list[str]:
    """Split long text into overlapping character windows at paragraph boundaries."""
    paragraphs = [part.strip() for part in text.split("\n\n") if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > CHUNK_SIZE:
            if current:
                chunks.append(current)
                current = ""
            start = 0
            while start < len(paragraph):
                end = min(start + CHUNK_SIZE, len(paragraph))
                chunks.append(paragraph[start:end].strip())
                if end == len(paragraph):
                    break
                start = max(end - CHUNK_OVERLAP, start + 1)
            continue
        candidate = f"{current}\n\n{paragraph}" if current else paragraph
        if len(candidate) <= CHUNK_SIZE:
            current = candidate
        else:
            chunks.append(current.strip())
            overlap = current[-CHUNK_OVERLAP:] if CHUNK_OVERLAP else ""
            current = f"{overlap}\n\n{paragraph}".strip()
    if current:
        chunks.append(current.strip())
    return [chunk for chunk in chunks if chunk]


def chunk_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Create traceable chunks while retaining source and page metadata."""
    chunks: list[dict[str, Any]] = []
    for record in records:
        for chunk_index, text in enumerate(_split_text(record["text"])):
            chunks.append({
                "chunk_id": len(chunks),
                "source": record["source"],
                "page": record["page"],
                "chunk_index": chunk_index,
                "text": text,
            })
    return chunks


def build_index(documents_dir: Path, vector_store_dir: Path, model_name: str) -> dict[str, Any]:
    """Extract all uploaded documents and rebuild a synchronized FAISS index."""
    vector_store_dir.mkdir(parents=True, exist_ok=True)
    files = [
        path
        for path in unique_paths(documents_dir.iterdir())
        if path.suffix.lower() in {".pdf", ".docx", ".txt"} and not path.name.startswith(".")
    ] if documents_dir.exists() else []
    all_chunks: list[dict[str, Any]] = []
    indexed_files: list[str] = []
    errors: list[str] = []
    for path in files:
        try:
            all_chunks.extend(chunk_records(extract_document(path)))
            indexed_files.append(path.name)
        except Exception as exc:
            errors.append(f"{path.name}: {exc}")
    if not all_chunks:
        raise ValueError("No usable document text was found. Upload a non-empty PDF, DOCX, or TXT file.")

    model = get_embedding_model(model_name)
    embeddings = model.encode(
        [chunk["text"] for chunk in all_chunks],
        normalize_embeddings=True,
        convert_to_numpy=True,
        show_progress_bar=False,
    ).astype("float32")
    index = faiss.IndexFlatIP(embeddings.shape[1])
    index.add(embeddings)
    manifest = {
        "embedding_model": model_name,
        "dimension": int(embeddings.shape[1]),
        "vector_count": int(index.ntotal),
        "documents": indexed_files,
    }
    with tempfile.TemporaryDirectory(dir=vector_store_dir.parent) as temp_dir:
        temp_path = Path(temp_dir)
        faiss.write_index(index, str(temp_path / "index.faiss"))
        (temp_path / "metadata.json").write_text(json.dumps(all_chunks, ensure_ascii=False, indent=2), encoding="utf-8")
        (temp_path / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        for filename in ("index.faiss", "metadata.json", "manifest.json"):
            os.replace(temp_path / filename, vector_store_dir / filename)
    return {"files": indexed_files, "chunks": len(all_chunks), "errors": errors}
